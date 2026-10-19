# -*- coding: utf-8 -*-
"""S7 汇编：把各话 draft.md 合成书，导出 md / docx / epub / mobi，并生成 PNG 封面。"""
from __future__ import annotations

import csv
import glob
import os
import re
import sys

import epub_writer as epub
import mobi_writer as mobi

CJK = re.compile(r"[\u4e00-\u9fff]")


def read_text(path: str) -> str:
    with open(path, encoding="utf-8") as fh:
        return fh.read()


def split_front_matter(text: str) -> tuple[dict, str]:
    if not text.startswith("---"):
        return {}, text
    parts = text.split("---", 2)
    if len(parts) < 3:
        return {}, text
    meta: dict = {}
    for line in parts[1].splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key, value = key.strip(), value.strip()
        if value.startswith("[") and value.endswith("]"):
            meta[key] = [v.strip() for v in value[1:-1].split(",") if v.strip()]
        else:
            meta[key] = value
    return meta, parts[2]


def chapter_key(meta: dict, path: str) -> float:
    for field in ("话", "chapter"):
        if field in meta:
            try:
                return float(str(meta[field]).strip())
            except ValueError:
                pass
    m = re.search(r"(\d+(?:\.\d+)?)", os.path.basename(path))
    return float(m.group(1)) if m else 9999.0


def count_chars(text: str) -> int:
    return len(CJK.findall(text))


def collect(work: str) -> list[dict]:
    paths = glob.glob(os.path.join(work, "ch*", "draft.md"))
    if not paths:
        paths = glob.glob(os.path.join(work, "*.md"))
    items = []
    for path in paths:
        meta, body = split_front_matter(read_text(path))
        title = meta.get("标题") or meta.get("title") or os.path.basename(os.path.dirname(path))
        items.append({"path": path, "meta": meta, "body": body.strip(),
                      "title": title, "key": chapter_key(meta, path)})
    items.sort(key=lambda x: x["key"])
    return items


def markdown_book(items: list[dict], title: str) -> str:
    lines = [f"# {title}", "", "## 目录", ""]
    for item in items:
        lines.append(f"- 第 {item['meta'].get('话', item['key'])} 话 · {item['title']}")
    lines += ["", ""]
    for item in items:
        body = re.sub(r"^#\s+(.*)$", r"## \1", item["body"], count=1, flags=re.M)
        lines += [body.strip(), "", "---", ""]
    return "\n".join(lines)


def chapter_paragraphs(item: dict) -> list[str]:
    body = re.sub(r"^#\s+.*$", "", item["body"], count=1, flags=re.M)
    return [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]


def write_docx(items: list[dict], out_path: str, title: str) -> str | None:
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        return None
    doc = docx.Document()
    doc.add_heading(title, level=0)
    doc.add_heading("目录", level=1)
    for it in items:
        doc.add_paragraph(f"第 {it['meta'].get('话', it['key'])} 话 · {it['title']}")
    for it in items:
        doc.add_page_break()
        doc.add_heading(it["title"], level=1)
        for para in chapter_paragraphs(it):
            if para.strip("◇ "):
                p = doc.add_paragraph(para)
                p.paragraph_format.first_line_indent = Pt(24)
    doc.save(out_path)
    return out_path


def write_epub(items: list[dict], out_path: str, title: str, cover: str | None) -> str:
    chapters = []
    for it in items:
        html = []
        for para in chapter_paragraphs(it):
            if para.strip("◇ "):
                html.append(f"<p>{_esc(para)}</p>")
            else:
                html.append("<p class='brk'>◇</p>")
        chapters.append({"title": f"第{it['meta'].get('话', it['key'])}话 {it['title']}",
                         "html": "".join(html)})
    return epub.build_epub(chapters, out_path, title, cover_png=cover)


def write_mobi(items: list[dict], out_path: str, title: str) -> str:
    chapters = []
    for it in items:
        chapters.append({"title": f"第{it['meta'].get('话', it['key'])}话 {it['title']}",
                         "text": "\n".join(chapter_paragraphs(it))})
    return mobi.build_mobi(chapters, out_path, title, author="")


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def make_cover(work: str, out_path: str, title: str) -> str | None:
    """优先用第一话首页做封面；没有页图时画一张纯色标题卡。"""
    candidates = sorted(glob.glob(os.path.join(work, "ch001", "pages", "p0001.*")))
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None
    if candidates:
        img = Image.open(candidates[0]).convert("RGB")
        img.thumbnail((1200, 1800))
        img.save(out_path, "PNG")
        return out_path
    img = Image.new("RGB", (900, 1200), (28, 30, 34))
    draw = ImageDraw.Draw(img)
    font = None
    for path in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"):
        if os.path.isfile(path):
            try:
                font = ImageFont.truetype(path, 64)
                break
            except OSError:
                continue
    draw.text((80, 520), title, fill=(240, 236, 228), font=font)
    img.save(out_path, "PNG")
    return out_path


def write_tables(items: list[dict], out_dir: str) -> None:
    with open(os.path.join(out_dir, "stats.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["chapter", "title", "words", "scenes", "characters"])
        for it in items:
            meta = it["meta"]
            w.writerow([meta.get("话", it["key"]), it["title"], count_chars(it["body"]),
                        len(meta.get("场景") or []), len(meta.get("人物") or [])])
    with open(os.path.join(out_dir, "index.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["chapter", "type", "value"])
        for it in items:
            meta = it["meta"]
            ch = meta.get("话", it["key"])
            for s in meta.get("场景") or []:
                w.writerow([ch, "scene", s])
            for c in meta.get("人物") or []:
                w.writerow([ch, "character", c])


def run(work: str, out: str, title: str = "漫画小说化",
        formats: list[str] | None = None, cover: bool = True) -> dict:
    items = collect(work)
    if not items:
        return {"error": f"no draft.md under {work}"}
    formats = formats or ["md", "docx", "epub", "mobi"]
    os.makedirs(out, exist_ok=True)
    result: dict = {"chapters": len(items),
                    "words": sum(count_chars(it["body"]) for it in items),
                    "files": {}}

    cover_path = None
    if cover:
        cover_path = make_cover(work, os.path.join(out, "cover.png"), title)
        if cover_path:
            result["files"]["cover"] = cover_path

    if "md" in formats:
        path = os.path.join(out, "book.md")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(markdown_book(items, title) + "\n")
        result["files"]["md"] = path
    if "docx" in formats:
        path = write_docx(items, os.path.join(out, "book.docx"), title)
        result["files"]["docx"] = path or "skipped (pip install python-docx)"
    if "epub" in formats:
        result["files"]["epub"] = write_epub(items, os.path.join(out, "book.epub"), title, cover_path)
    if "mobi" in formats:
        path = write_mobi(items, os.path.join(out, "book.mobi"), title)
        result["files"]["mobi"] = path
        result["mobi_check"] = mobi.inspect(path)

    write_tables(items, out)
    result["files"]["stats"] = os.path.join(out, "stats.csv")
    result["files"]["index"] = os.path.join(out, "index.csv")
    return result


if __name__ == "__main__":
    import argparse
    import json
    ap = argparse.ArgumentParser(description="汇编成书")
    ap.add_argument("--work", default=".")
    ap.add_argument("--out", default="build")
    ap.add_argument("--title", default="漫画小说化")
    ap.add_argument("--formats", nargs="*", default=["md", "docx", "epub", "mobi"])
    ap.add_argument("--no-cover", action="store_true")
    args = ap.parse_args()
    print(json.dumps(run(args.work, args.out, args.title, args.formats, not args.no_cover),
                     ensure_ascii=False, indent=1))
    sys.exit(0)
